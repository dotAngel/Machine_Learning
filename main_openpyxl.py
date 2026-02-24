from openpyxl import load_workbook
import os
from calcular_m import calcular_pendiente
from calcular_b import calcular_intercepto

# Leer datos desde XLSX
script_dir = os.path.dirname(os.path.abspath(__file__))
xlsx_path = os.path.join(script_dir, 'data_openpyxl.xlsx')
wb = load_workbook(xlsx_path)
ws = wb.active

x = []
y = []
predecir = 15

for row in ws.iter_rows(min_row=2, values_only=True):
    if row[0] is not None and row[1] is not None:
        x.append(row[0])
        y.append(row[1])

# Calcular m y b
n = len(x)
sum_x = sum(x)
sum_y = sum(y)
sum_xy = sum(xi * yi for xi, yi in zip(x, y))
sum_x2 = sum(xi**2 for xi in x)

m = calcular_pendiente(n, sum_x, sum_y, sum_xy, sum_x2)
b = calcular_intercepto(n, sum_x, sum_y, m)

# Función para predecir
def pred(x_val):
    return m * x_val + b

def sf(op):
    if op:
        print("Valor de m =", m)
        print("Valor de b =", b)
    else:
        print(f"Ecuacion de la recta: y = {m} X  {f'- {abs(b)}' if b < 0 else f'+ {b}'}")

print(f"Valor para la prediccion de x = {predecir} es y = {pred(predecir)}")
sf(0)
